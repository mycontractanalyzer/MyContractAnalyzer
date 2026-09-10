"""MyContractAnalyzer API — шаг 1 переезда. Работает параллельно со Streamlit."""
import json

from fastapi import Depends, FastAPI, File, Header, HTTPException, UploadFile
from fastapi.responses import Response, StreamingResponse
from pydantic import BaseModel

import secrets as py_secrets

from core.analyzer import analyze_contract_stream, split_report_highlights
from core.contracts import save_analysis, save_contract, spend_checks
from core.mailer import send_code_email
from core.records import save_highlights
from database.connection import get_connection
from utils.auth import hash_password, is_valid_email, verify_password

app = FastAPI(title="MyContractAnalyzer API", version="1.0")


class RegisterIn(BaseModel):
    email: str
    password: str
    password2: str


class VerifyIn(BaseModel):
    email: str
    code: str


class LoginIn(BaseModel):
    email: str
    password: str


class AnalyzeIn(BaseModel):
    text: str
    depth: str = "standard"
    contract_type: str = ""
    role: str = ""
    comment: str = ""


def _user_by_token(token: str):
    if not token:
        return None
    conn = get_connection()
    row = conn.execute("SELECT * FROM users WHERE token = ?", (token,)).fetchone()
    conn.close()
    return dict(row) if row else None


def _auth(authorization: str = Header(default="")):
    user = _user_by_token(authorization.replace("Bearer ", "").strip())
    if not user:
        raise HTTPException(401, "Не авторизован")
    return user


@app.get("/api/health")
def health():
    return {"ok": True, "service": "mca-api"}


@app.post("/api/register")
def register(data: RegisterIn):
    email = data.email.strip().lower()
    if not is_valid_email(email):
        raise HTTPException(400, "Некорректный email")
    if len(data.password or "") < 6:
        raise HTTPException(400, "Пароль должен быть не короче 6 символов")
    if data.password != data.password2:
        raise HTTPException(400, "Пароли не совпадают")
    conn = get_connection()
    cols = [r["name"] for r in conn.execute("PRAGMA table_info(users)").fetchall()]
    if "verified" not in cols:
        conn.execute("ALTER TABLE users ADD COLUMN verified INTEGER DEFAULT 0")
        conn.execute("ALTER TABLE users ADD COLUMN verify_code TEXT DEFAULT ''")
        conn.commit()
    try:
        conn.execute("INSERT INTO users (email, password_hash) VALUES (?, ?)",
                     (email, hash_password(data.password)))
        conn.commit()
    except Exception:
        conn.close()
        raise HTTPException(409, "Пользователь с таким email уже существует")
    code = str(py_secrets.randbelow(900000) + 100000)
    conn.execute("UPDATE users SET verified = 0, verify_code = ? WHERE email = ?", (code, email))
    conn.commit()
    conn.close()
    send_code_email(email, code)
    return {"ok": True, "message": "Код подтверждения отправлен на почту"}


@app.post("/api/verify")
def verify(data: VerifyIn):
    email = data.email.strip().lower()
    conn = get_connection()
    row = conn.execute("SELECT * FROM users WHERE email = ?", (email,)).fetchone()
    if row and (row["verify_code"] or "") == data.code.strip():
        conn.execute("UPDATE users SET verified = 1, verify_code = '' WHERE email = ?", (email,))
        conn.commit()
        conn.close()
        return {"ok": True}
    conn.close()
    raise HTTPException(400, "Неверный код")


@app.post("/api/login")
def login(data: LoginIn):
    email = data.email.strip().lower()
    conn = get_connection()
    row = conn.execute("SELECT * FROM users WHERE email = ?", (email,)).fetchone()
    if row is None or not verify_password(data.password, row["password_hash"]):
        conn.close()
        raise HTTPException(401, "Неверный email или пароль")
    if not row["verified"]:
        conn.close()
        raise HTTPException(403, "Почта не подтверждена")
    token = py_secrets.token_hex(16)
    conn.execute("UPDATE users SET token = ? WHERE id = ?", (token, row["id"]))
    conn.commit()
    conn.close()
    return {"ok": True, "token": token}


@app.get("/api/me")
def me(user=Depends(_auth)):
    return {"email": user["email"], "tariff": user["tariff"], "checks_left": user["checks_left"]}


@app.post("/api/analyze")
def analyze(data: AnalyzeIn, user=Depends(_auth)):
    if user["checks_left"] < 1:
        raise HTTPException(402, "Недостаточно проверок")
    cap = TIER_CHARS.get(user["tariff"], 15000)
    if len(data.text) > cap:
        raise HTTPException(413, f"Лимит тарифа {user['tariff']}: до {cap} символов за одну проверку. Раздели документ или повысь тариф.")
    import re as _re
    pre = max(len(_re.findall(r"заключили настоящий договор", data.text, _re.I)),
              len(_re.findall(r"именуем\w* в дальнейшем [«\"']?Арендодатель", data.text, _re.I)))
    if pre > 1:
        raise HTTPException(422, f"Похоже, в тексте несколько договоров ({pre}). Загрузите один договор: 1 проверка = 1 документ.")
    cap = TIER_CHARS.get(user["tariff"], 15000)
    if len(data.text) > cap:
        raise HTTPException(413, f"Лимит тарифа {user['tariff']}: до {cap} символов за одну проверку. Раздели документ или повысь тариф.")
    gen, model = analyze_contract_stream(
        data.text, user["tariff"], data.contract_type, data.role, data.comment, depth=data.depth)

    def stream():
        chunks = []
        for ch in gen:
            chunks.append(ch)
            yield f"data: {json.dumps({'chunk': ch}, ensure_ascii=False)}\n\n"
        report, hl = split_report_highlights("".join(chunks))
        spend_checks(user["id"], len(data.text))
        cid = save_contract(user["id"], data.contract_type or "Договор", data.role, data.text)
        aid = save_analysis(user["id"], cid, model, report)
        if hl:
            save_highlights(aid, hl)
        yield f"data: {json.dumps({'done': True, 'analysis_id': aid}, ensure_ascii=False)}\n\n"

    return StreamingResponse(stream(), media_type="text/event-stream")


@app.get("/api/analyses")
def analyses(user=Depends(_auth)):
    conn = get_connection()
    rows = conn.execute(
        "SELECT a.id, a.created_at, a.report, c.contract_type "
        "FROM analyses a JOIN contracts c ON c.id = a.contract_id "
        "WHERE a.user_id = ? ORDER BY a.id DESC LIMIT 30", (user["id"],)).fetchall()
    conn.close()
    return [{"id": r["id"], "created_at": r["created_at"], "type": r["contract_type"],
             "preview": (r["report"] or "")[:120]} for r in rows]


@app.get("/api/analyses/{aid}")
def analysis_detail(aid: int, user=Depends(_auth)):
    conn = get_connection()
    row = conn.execute(
        "SELECT * FROM analyses WHERE id = ? AND user_id = ?", (aid, user["id"])).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "Отчёт не найден")
    return {"id": row["id"], "report": row["report"],
            "highlights": row["highlights"], "created_at": row["created_at"]}


class _FileShim:
    """Обёртка, чтобы переиспользовать ридер файлов из Streamlit-версии."""
    def __init__(self, name, mime, data):
        self.name = name
        self.type = mime
        self._data = data

    def read(self):
        return self._data

    def getvalue(self):
        return self._data


@app.post("/api/upload")
async def upload(file: UploadFile = File(...), user=Depends(_auth)):
    data = await file.read()
    name = (file.filename or "").lower()
    text = ""
    try:
        from core.file_reader import read_uploaded_file
        text = read_uploaded_file(_FileShim(file.filename, file.content_type or "", data)) or ""
    except Exception:
        text = ""
    if not text.strip():
        import io
        if name.endswith((".txt", ".md")):
            text = data.decode("utf-8", errors="ignore")
        elif name.endswith(".pdf"):
            try:
                import pdfplumber
                with pdfplumber.open(io.BytesIO(data)) as pdf:
                    text = "\n".join((p.extract_text() or "") for p in pdf.pages)
            except Exception:
                try:
                    from pypdf import PdfReader
                except Exception:
                    from PyPDF2 import PdfReader
                reader = PdfReader(io.BytesIO(data))
                text = "\n".join((p.extract_text() or "") for p in reader.pages)
        elif name.endswith(".docx"):
            import docx
            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" ".join(c.text for c in row.cells))
            text = "\n".join(parts)
        else:
            raise HTTPException(400, "Форматы: TXT, MD, PDF, DOCX")
    if not text.strip():
        raise HTTPException(400, "В файле нет текста (возможно, это скан — используй старую версию с OCR)")
    return {"ok": True, "text": text, "chars": len(text)}


@app.get("/api/analyses/{aid}/pdf")
def analysis_pdf(aid: int, user=Depends(_auth)):
    conn = get_connection()
    row = conn.execute("SELECT report FROM analyses WHERE id = ? AND user_id = ?",
                       (aid, user["id"])).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "Отчёт не найден")
    from storage.pdf_generator import generate_report_pdf
    data = generate_report_pdf(row["report"], user["email"])
    return Response(content=data, media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="report_{aid}.pdf"'})


@app.get("/api/analyses/{aid}/docx")
def analysis_docx(aid: int, user=Depends(_auth)):
    conn = get_connection()
    row = conn.execute("SELECT report FROM analyses WHERE id = ? AND user_id = ?",
                       (aid, user["id"])).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "Отчёт не найден")
    from storage.docx_generator import generate_report_docx
    data = generate_report_docx(row["report"], user["email"])
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="report_{aid}.docx"'})


@app.get("/api/analyses/{aid}/contract")
def analysis_contract(aid: int, user=Depends(_auth)):
    conn = get_connection()
    a = conn.execute("SELECT contract_id FROM analyses WHERE id = ? AND user_id = ?",
                     (aid, user["id"])).fetchone()
    if not a:
        conn.close()
        raise HTTPException(404, "Отчёт не найден")
    c = conn.execute("SELECT source_text, contract_type FROM contracts WHERE id = ?",
                     (a["contract_id"],)).fetchone()
    conn.close()
    return {"text": c["source_text"], "type": c["contract_type"]}


class ToolIn(BaseModel):
    analysis_id: int
    scenario: str = ""


@app.post("/api/tools/{tool}")
def run_tool(tool: str, data: ToolIn, user=Depends(_auth)):
    conn = get_connection()
    a = conn.execute("SELECT * FROM analyses WHERE id = ? AND user_id = ?",
                     (data.analysis_id, user["id"])).fetchone()
    if not a:
        conn.close()
        raise HTTPException(404, "Отчёт не найден")
    c = conn.execute("SELECT * FROM contracts WHERE id = ?", (a["contract_id"],)).fetchone()
    conn.close()
    a, c = dict(a), dict(c)
    text, report, tariff = c["source_text"], a["report"], user["tariff"]
    from core.analyzer import (generate_benchmark, generate_letter, generate_missing,
                               generate_negotiation, generate_passport, generate_redline,
                               generate_whatif, translate_contract)
    from core.extra_ai import generate_precedent
    mp = {
        "redline": lambda: generate_redline(text, report, tariff),
        "letter": lambda: generate_letter(text, report, tariff, c["contract_type"] or "", c["role"] or ""),
        "negotiation": lambda: generate_negotiation(text, report, tariff),
        "whatif": lambda: generate_whatif(text, report, data.scenario or "Своя ситуация", tariff),
        "benchmark": lambda: generate_benchmark(text, report, tariff),
        "passport": lambda: generate_passport(text),
        "missing": lambda: generate_missing(text, report, tariff),
        "translate": lambda: translate_contract(text),
        "precedent": lambda: generate_precedent(text, report, tariff),
    }
    fn = mp.get(tool)
    if not fn:
        raise HTTPException(400, "Неизвестный инструмент")
    return {"ok": True, "text": fn()}


@app.get("/api/analyses/{aid}/protocol")
def analysis_protocol(aid: int, user=Depends(_auth)):
    conn = get_connection()
    a = conn.execute("SELECT * FROM analyses WHERE id = ? AND user_id = ?",
                     (aid, user["id"])).fetchone()
    if not a:
        conn.close()
        raise HTTPException(404, "Отчёт не найден")
    c = conn.execute("SELECT * FROM contracts WHERE id = ?", (a["contract_id"],)).fetchone()
    conn.close()
    a, c = dict(a), dict(c)
    from core.protocol import generate_protocol, protocol_docx
    rows = generate_protocol(a["report"], c["source_text"], user["tariff"])
    data = protocol_docx(rows, user["email"], a.get("title") or "Договор")
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="protocol_{aid}.docx"'})


class LawyerIn(BaseModel):
    question: str
    history: list = []


LAWYER_LIMITS = {"Pro": 25, "Business": 15, "Business Pro": 50}
_LAWYER_USAGE = {}


@app.post("/api/lawyer")
def lawyer(data: LawyerIn, user=Depends(_auth)):
    from datetime import date
    limit = LAWYER_LIMITS.get(user["tariff"], 0)
    if limit == 0:
        raise HTTPException(403, "ИИ-юрист 24/7 доступен на тарифах Pro, Business и Business Pro")
    key = f"{user['email']}|{date.today().isoformat()}"
    used = _lawyer_used(user["email"])
    if used >= limit:
        raise HTTPException(429, f"Дневной лимит исчерпан: {limit} обращений на тарифе {user['tariff']}. Счётчик сбросится в полночь.")
    _lawyer_inc(user["email"])
    from core.analyzer import lawyer247_stream
    gen, model = lawyer247_stream("Отвечай строго на русском языке, без иноязычных вставок. " + data.question, data.history, user["tariff"])

    def stream():
        out = []
        for ch in gen:
            out.append(ch)
            yield f"data: {json.dumps({'chunk': ch}, ensure_ascii=False)}\n\n"
        try:
            conn = get_connection()
            conn.execute("CREATE TABLE IF NOT EXISTS lawyer_chats (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, analysis_id INTEGER, question TEXT, answer TEXT, created_at TEXT DEFAULT (datetime('now')))")
            conn.execute("INSERT INTO lawyer_chats (user_id, analysis_id, question, answer) VALUES (?,?,?,?)",
                         (user["id"], None, data.question, "".join(out)))
            conn.commit()
            conn.close()
        except Exception:
            pass
        yield f"data: {json.dumps({'done': True, 'left': limit - used - 1}, ensure_ascii=False)}\n\n"

    return StreamingResponse(stream(), media_type="text/event-stream")


class FeedbackIn(BaseModel):
    analysis_id: int
    rating: int
    comment: str = ""


@app.post("/api/feedback")
def feedback(data: FeedbackIn, user=Depends(_auth)):
    from core.feedback import upsert_feedback
    upsert_feedback(data.analysis_id, user["id"], data.rating, data.comment)
    return {"ok": True}


class SupportIn(BaseModel):
    topic: str
    message: str


@app.post("/api/support")
def support(data: SupportIn, user=Depends(_auth)):
    conn = get_connection()
    conn.execute("""CREATE TABLE IF NOT EXISTS support_messages (
        id INTEGER PRIMARY KEY AUTOINCREMENT, email TEXT, topic TEXT, message TEXT,
        replied INTEGER DEFAULT 0, created_at TEXT DEFAULT (datetime('now')))""")
    conn.execute("INSERT INTO support_messages (email, topic, message) VALUES (?,?,?)",
                 (user["email"], data.topic, data.message))
    conn.commit()
    conn.close()
    return {"ok": True}


class PassIn(BaseModel):
    old: str
    new: str


@app.post("/api/change_password")
def change_pass(data: PassIn, user=Depends(_auth)):
    from utils.auth import change_password
    ok, msg = change_password(user["id"], data.old, data.new)
    if not ok:
        raise HTTPException(400, msg)
    return {"ok": True, "message": msg}
    

class TitleIn(BaseModel):
    title: str


@app.post("/api/analyses/{aid}/share")
def share_analysis(aid: int, user=Depends(_auth)):
    conn = get_connection()
    conn.execute("UPDATE analyses SET share = 1 WHERE id = ? AND user_id = ?",
                 (aid, user["id"]))
    conn.commit()
    conn.close()
    return {"ok": True, "url": f"http://185.171.82.207/report.html?id={aid}"}


@app.post("/api/analyses/{aid}/rename")
def rename_analysis_api(aid: int, data: TitleIn, user=Depends(_auth)):
    conn = get_connection()
    conn.execute("UPDATE analyses SET title = ? WHERE id = ? AND user_id = ?",
                 (data.title, aid, user["id"]))
    conn.commit()
    conn.close()
    return {"ok": True}


@app.post("/api/analyses/{aid}/email")
def email_analysis(aid: int, user=Depends(_auth)):
    conn = get_connection()
    row = conn.execute("SELECT report, title FROM analyses WHERE id = ? AND user_id = ?",
                       (aid, user["id"])).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "Отчёт не найден")
    from storage.pdf_generator import generate_report_pdf
    from core.mailer import send_report_email
    pdf = generate_report_pdf(row["report"], user["email"])
    ok = send_report_email(user["email"], pdf, row["title"] or "Договор")
    if not ok:
        raise HTTPException(500, "Не удалось отправить письмо")
    return {"ok": True}


@app.get("/api/public/analyses/{aid}")
def public_analysis(aid: int):
    conn = get_connection()
    row = conn.execute(
        "SELECT a.report, a.highlights, a.title, c.contract_type "
        "FROM analyses a JOIN contracts c ON c.id = a.contract_id "
        "WHERE a.id = ? AND a.share = 1", (aid,)).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "Отчёт недоступен или скрыт владельцем")
    return dict(row)


@app.post("/api/upload_ocr")
async def upload_ocr(file: UploadFile = File(...), user=Depends(_auth)):
    data = await file.read()
    from core.vision import ocr_image
    try:
        text = ocr_image(data, file.content_type or "image/png")
    except Exception:
        raise HTTPException(400, "Не удалось распознать фото. Попробуй более чёткий снимок.")
    if not (text or "").strip():
        raise HTTPException(400, "В фото не найдено текста")
    return {"ok": True, "text": text, "chars": len(text)}


def _send_text_email(to: str, subject: str, body: str):
    import smtplib
    import ssl
    from email.mime.text import MIMEText
    from email.header import Header
    import streamlit as st
    sender = st.secrets.get("GMAIL_EMAIL", "") or getattr(config, "GMAIL_EMAIL", "")
    pwd = st.secrets.get("GMAIL_APP_PASSWORD", "") or getattr(config, "GMAIL_APP_PASSWORD", "")
    if not sender or not pwd:
        return False
    msg = MIMEText(body, "plain", "utf-8")
    msg["Subject"] = Header(subject, "utf-8")
    msg["From"] = sender
    msg["To"] = to
    with smtplib.SMTP_SSL("smtp.gmail.com", 465, context=ssl.create_default_context()) as s:
        s.login(sender, pwd)
        s.sendmail(sender, [to], msg.as_string())
    return True


def _admin(user):
    if user["email"] not in ["mycontractanalyzer@gmail.com"]:
        raise HTTPException(403, "Доступ только для администратора")


def _col(conn, table, name):
    return name in [r["name"] for r in conn.execute(f"PRAGMA table_info({table})")]


@app.get("/api/admin/stats")
def admin_stats(user=Depends(_auth)):
    _admin(user)
    conn = get_connection()
    q = lambda s: conn.execute(s).fetchone()[0]
    stats = {
        "users": q("SELECT COUNT(*) FROM users"),
        "verified": q("SELECT COUNT(*) FROM users WHERE verified = 1"),
        "analyses": q("SELECT COUNT(*) FROM analyses"),
        "contracts": q("SELECT COUNT(*) FROM contracts"),
        "feedback_avg": q("SELECT ROUND(AVG(rating),2) FROM feedback"),
        "support_open": q("SELECT COUNT(*) FROM support_messages WHERE replied = 0"),
        "laws": q("SELECT COUNT(*) FROM laws"),
    }
    if _col(conn, "users", "created_at"):
        stats["reg_today"] = q("SELECT COUNT(*) FROM users WHERE date(created_at) = date('now')")
    stats["recent_users"] = [dict(r) for r in conn.execute(
        "SELECT email, tariff, checks_left FROM users ORDER BY id DESC LIMIT 15")]
    conn.close()
    return stats


@app.get("/api/admin/support")
def admin_support(user=Depends(_auth)):
    _admin(user)
    conn = get_connection()
    conn.execute("""CREATE TABLE IF NOT EXISTS support_messages (
        id INTEGER PRIMARY KEY AUTOINCREMENT, email TEXT, topic TEXT, message TEXT,
        replied INTEGER DEFAULT 0, created_at TEXT DEFAULT (datetime('now')))""")
    rows = [dict(r) for r in conn.execute(
        "SELECT * FROM support_messages ORDER BY id DESC LIMIT 50")]
    conn.close()
    return rows


class ReplyIn(BaseModel):
    answer: str


@app.post("/api/admin/support/{mid}/reply")
def admin_reply(mid: int, data: ReplyIn, user=Depends(_auth)):
    _admin(user)
    conn = get_connection()
    row = conn.execute("SELECT * FROM support_messages WHERE id = ?", (mid,)).fetchone()
    if not row:
        conn.close()
        raise HTTPException(404, "Обращение не найдено")
    ok = _send_text_email(
        row["email"],
        f"MyContractAnalyzer · Ответ поддержки по теме «{row['topic']}»",
        f"Здравствуйте!\n\nВаше обращение: {row['message']}\n\nОтвет: {data.answer}\n\nС уважением, команда MyContractAnalyzer")
    if ok:
        conn.execute("UPDATE support_messages SET replied = 1 WHERE id = ?", (mid,))
        conn.commit()
    conn.close()
    if not ok:
        raise HTTPException(500, "Почта не настроена")
    return {"ok": True}


_LAWS_JOB = {"running": False, "done": 0, "total": 0, "last": "", "error": ""}


@app.get("/api/admin/laws")
def admin_laws(user=Depends(_auth)):
    _admin(user)
    conn = get_connection()
    rows = [dict(r) for r in conn.execute(
        "SELECT code, title, LENGTH(COALESCE(full_text,'')) AS ft FROM laws ORDER BY code")]
    conn.close()
    return {"laws": rows, "job": _LAWS_JOB}


@app.post("/api/admin/laws/reload")
def admin_laws_reload(user=Depends(_auth)):
    _admin(user)
    import threading
    if _LAWS_JOB["running"]:
        return {"ok": False, "detail": "Уже выполняется"}

    def run():
        from core.laws_autoload import DEFAULT_PACK, autoload_law
        _LAWS_JOB.update(running=True, done=0, total=len(DEFAULT_PACK), last="", error="")
        try:
            for prefix, title, cands in DEFAULT_PACK:
                n, err, source = autoload_law(prefix, title, cands)
                _LAWS_JOB["done"] += 1
                _LAWS_JOB["last"] = f"{title}: статей {n}" if n else f"{title}: {err}"
        except Exception as e:
            _LAWS_JOB["error"] = str(e)
        finally:
            _LAWS_JOB["running"] = False

    threading.Thread(target=run, daemon=True).start()
    return {"ok": True}


@app.get("/api/admin/users")
def admin_users(user=Depends(_auth)):
    _admin(user)
    conn = get_connection()
    rows = [dict(r) for r in conn.execute(
        "SELECT id, email, tariff, checks_left, verified, created_at FROM users ORDER BY id DESC")]
    conn.close()
    return rows


class UserTariffIn(BaseModel):
    tariff: str


@app.post("/api/admin/users/{uid}/tariff")
def admin_set_tariff(uid: int, data: UserTariffIn, user=Depends(_auth)):
    _admin(user)
    conn = get_connection()
    conn.execute("UPDATE users SET tariff = ? WHERE id = ?", (data.tariff, uid))
    conn.commit()
    conn.close()
    return {"ok": True}


class UserChecksIn(BaseModel):
    delta: int


@app.post("/api/admin/users/{uid}/checks")
def admin_add_checks(uid: int, data: UserChecksIn, user=Depends(_auth)):
    _admin(user)
    conn = get_connection()
    conn.execute("UPDATE users SET checks_left = MAX(0, checks_left + ?) WHERE id = ?",
                 (data.delta, uid))
    conn.commit()
    conn.close()
    return {"ok": True}


@app.get("/api/admin/feedback")
def admin_feedback(user=Depends(_auth)):
    _admin(user)
    conn = get_connection()
    rows = [dict(r) for r in conn.execute(
        "SELECT f.*, u.email FROM feedback f LEFT JOIN users u ON u.id = f.user_id "
        "ORDER BY f.id DESC LIMIT 100")]
    conn.close()
    return rows


@app.get("/api/admin/consults")
def admin_consults(user=Depends(_auth)):
    _admin(user)
    conn = get_connection()
    try:
        rows = [dict(r) for r in conn.execute(
            "SELECT c.*, u.email FROM consult_requests c LEFT JOIN users u ON u.id = c.user_id "
            "ORDER BY c.id DESC LIMIT 100")]
    except Exception:
        rows = []
    conn.close()
    return rows


@app.get("/api/admin/series")
def admin_series(user=Depends(_auth)):
    _admin(user)
    conn = get_connection()
    regs = [dict(r) for r in conn.execute(
        "SELECT date(created_at) d, COUNT(*) c FROM users GROUP BY d ORDER BY d")]
    anls = [dict(r) for r in conn.execute(
        "SELECT date(created_at) d, COUNT(*) c FROM analyses GROUP BY d ORDER BY d")]
    paid = conn.execute("SELECT COUNT(*) c FROM users WHERE tariff != 'Free'").fetchone()["c"]
    try:
        promos = [dict(r) for r in conn.execute(
            "SELECT code, used_count FROM promocodes ORDER BY used_count DESC LIMIT 5")]
    except Exception:
        promos = []
    conn.close()
    return {"regs": regs, "anals": anls, "paid": paid, "top_promos": promos}


class ResetPassIn(BaseModel):
    password: str


@app.post("/api/admin/users/{uid}/reset_password")
def admin_reset_pass(uid: int, data: ResetPassIn, user=Depends(_auth)):
    _admin(user)
    from utils.auth import reset_password_admin
    reset_password_admin(uid, data.password)
    return {"ok": True}


@app.post("/api/admin/users/{uid}/delete")
def admin_delete_user(uid: int, user=Depends(_auth)):
    _admin(user)
    from utils.auth import delete_user
    delete_user(uid)
    return {"ok": True}


class GrantIn(BaseModel):
    tariff: str
    promo_code: str = ""


@app.post("/api/admin/users/{uid}/grant")
def admin_grant(uid: int, data: GrantIn, user=Depends(_auth)):
    _admin(user)
    from core.tariffs import TARIFFS
    if data.tariff not in TARIFFS:
        raise HTTPException(400, "Неизвестный тариф")
    conn = get_connection()
    conn.execute("UPDATE users SET tariff = ?, checks_left = ? WHERE id = ?",
                 (data.tariff, TIER_CHECKS.get(data.tariff, 3), uid))
    conn.commit()
    conn.close()
    msg = ""
    if data.promo_code.strip():
        from core.promocodes import register_discount_use
        ok, m = register_discount_use(data.promo_code.strip(), uid)
        if not ok:
            msg = f"Промокод не засчитан: {m}"
    return {"ok": True, "message": msg}


class ChecksSetIn(BaseModel):
    value: int
    mode: str = "add"


@app.post("/api/admin/users/{uid}/checks_set")
def admin_checks_set(uid: int, data: ChecksSetIn, user=Depends(_auth)):
    _admin(user)
    conn = get_connection()
    if data.mode == "set":
        conn.execute("UPDATE users SET checks_left = ? WHERE id = ?", (max(0, data.value), uid))
    else:
        conn.execute("UPDATE users SET checks_left = MAX(0, checks_left + ?) WHERE id = ?",
                     (data.value, uid))
    conn.commit()
    conn.close()
    return {"ok": True}


@app.get("/api/admin/promocodes")
def admin_promos(user=Depends(_auth)):
    _admin(user)
    from core.promocodes import list_promocodes
    return list_promocodes()


class PromoCreateIn(BaseModel):
    kind: str
    checks_bonus: int = 0
    discount_rub: int = 0
    min_tariff: str = ""
    expires_at: str = ""
    custom_code: str = ""


@app.post("/api/admin/promocodes")
def admin_promo_create(data: PromoCreateIn, user=Depends(_auth)):
    _admin(user)
    from core.promocodes import create_promocode
    ok, result = create_promocode(
        kind=data.kind,
        value=data.checks_bonus if data.kind == "checks" else 0,
        discount_rub=data.discount_rub,
        min_tariff=data.min_tariff or None,
        checks_bonus=data.checks_bonus,
        expires_at=data.expires_at or None,
        custom_code=data.custom_code or None)
    if not ok:
        raise HTTPException(400, result)
    return {"ok": True, "code": result}


@app.post("/api/admin/promocodes/{code}/deactivate")
def admin_promo_off(code: str, user=Depends(_auth)):
    _admin(user)
    from core.promocodes import deactivate_promocode
    deactivate_promocode(code)
    return {"ok": True}


@app.post("/api/admin/promocodes/{code}/delete")
def admin_promo_del(code: str, user=Depends(_auth)):
    _admin(user)
    conn = get_connection()
    conn.execute("DELETE FROM promocodes WHERE code = ?", (code,))
    conn.commit()
    conn.close()
    return {"ok": True}


class ResetReqIn(BaseModel):
    email: str


@app.post("/api/reset_request")
def reset_request(data: ResetReqIn):
    import random
    email = data.email.strip().lower()
    conn = get_connection()
    cols = [r["name"] for r in conn.execute("PRAGMA table_info(users)").fetchall()]
    if "reset_code" not in cols:
        conn.execute("ALTER TABLE users ADD COLUMN reset_code TEXT DEFAULT ''")
        conn.commit()
    row = conn.execute("SELECT id FROM users WHERE email = ?", (email,)).fetchone()
    if not row:
        conn.close()
        return {"ok": True}
    code = str(random.randint(100000, 999999))
    conn.execute("UPDATE users SET reset_code = ? WHERE id = ?", (code, row["id"]))
    conn.commit()
    conn.close()
    _send_text_email(email, "MyContractAnalyzer — восстановление пароля",
                     f"Ваш код для смены пароля: {code}\nЕсли вы не запрашивали сброс — проигнорируйте письмо.")
    return {"ok": True}


class ResetConfirmIn(BaseModel):
    email: str
    code: str
    new_password: str


@app.post("/api/reset_confirm")
def reset_confirm(data: ResetConfirmIn):
    email = data.email.strip().lower()
    conn = get_connection()
    row = conn.execute("SELECT * FROM users WHERE email = ?", (email,)).fetchone()
    ok = bool(row) and (row["reset_code"] or "") == data.code.strip() and len(data.new_password) >= 6
    if ok:
        conn.execute("UPDATE users SET password_hash = ?, reset_code = '' WHERE id = ?",
                     (hash_password(data.new_password), row["id"]))
        conn.commit()
    conn.close()
    if not ok:
        raise HTTPException(400, "Неверный код или пароль короче 6 символов")
    return {"ok": True}


TIER_CHECKS = {"Free": 3, "Light": 35, "Standard": 65, "Pro": 125,
               "Business": 150, "Business Pro": 200}

TIER_CHARS = {"Free": 15000, "Light": 35000, "Standard": 50000, "Pro": 100000,
              "Business": 150000, "Business Pro": 200000}


EXTENDED_PACK = [
    ("ЖК", "Жилищный кодекс РФ", ["Жилищный кодекс Российской Федерации"]),
    ("ЗК", "Земельный кодекс РФ", ["Земельный кодекс Российской Федерации"]),
    ("НК-1", "Налоговый кодекс РФ (часть 1)", ["Налоговый кодекс Российской Федерации (часть первая)"]),
    ("НК-2", "Налоговый кодекс РФ (часть 2)", ["Налоговый кодекс Российской Федерации (часть вторая)"]),
    ("353-ФЗ", "ФЗ № 353-ФЗ О потребительском кредите (займе)", ["Федеральный закон № 353-ФЗ О потребительском кредите (займе)"]),
    ("214-ФЗ", "ФЗ № 214-ФЗ Об участии в долевом строительстве", ["Федеральный закон № 214-ФЗ Об участии в долевом строительстве многоквартирных домов и иных объектов недвижимости"]),
    ("127-ФЗ", "ФЗ № 127-ФЗ О несостоятельности (банкротстве)", ["Федеральный закон № 127-ФЗ О несостоятельности (банкротстве)"]),
    ("63-ФЗ", "ФЗ № 63-ФЗ Об электронной подписи", ["Федеральный закон № 63-ФЗ Об электронной подписи"]),
    ("98-ФЗ", "ФЗ № 98-ФЗ О коммерческой тайне", ["Федеральный закон № 98-ФЗ О коммерческой тайне"]),
    ("38-ФЗ", "ФЗ № 38-ФЗ О рекламе", ["Федеральный закон № 38-ФЗ О рекламе"]),
    ("164-ФЗ", "ФЗ № 164-ФЗ О финансовой аренде (лизинге)", ["Федеральный закон № 164-ФЗ О финансовой аренде (лизинге)"]),
    ("229-ФЗ", "ФЗ № 229-ФЗ Об исполнительном производстве", ["Федеральный закон № 229-ФЗ Об исполнительном производстве"]),
    ("135-ФЗ", "ФЗ № 135-ФЗ О защите конкуренции", ["Федеральный закон № 135-ФЗ О защите конкуренции"]),
    ("ВоздК", "Воздушный кодекс РФ", ["Воздушный кодекс Российской Федерации"]),
    ("УАТ", "Устав автомобильного транспорта", ["Устав автомобильного транспорта и городского наземного электрического транспорта"]),
    ("ГК-1", "Гражданский кодекс РФ (часть первая)", ["Гражданский кодекс Российской Федерации (часть первая)"]),
    ("ГК-3", "Гражданский кодекс РФ (часть третья)", ["Гражданский кодекс Российской Федерации (часть третья)"]),
    ("ГК-4", "Гражданский кодекс РФ (часть четвёртая)", ["Гражданский кодекс Российской Федерации (часть четвертая)"]),
    ("УК", "Уголовный кодекс РФ", ["Уголовный кодекс Российской Федерации"]),
    ("ВК", "Водный кодекс РФ", ["Водный кодекс Российской Федерации"]),
    ("ЛК", "Лесной кодекс РФ", ["Лесной кодекс Российской Федерации"]),
    ("КТМ", "Кодекс торгового мореплавания РФ", ["Кодекс торгового мореплавания Российской Федерации"]),
    ("КВВТ", "Кодекс внутреннего водного транспорта РФ", ["Кодекс внутреннего водного транспорта Российской Федерации"]),
    ("УЖТ", "Устав железнодорожного транспорта РФ", ["Устав железнодорожного транспорта Российской Федерации"]),
    ("44-ФЗ", "ФЗ № 44-ФЗ О контрактной системе в сфере закупок", ["Федеральный закон № 44-ФЗ О контрактной системе в сфере закупок товаров, работ, услуг для обеспечения государственных и муниципальных нужд"]),
    ("223-ФЗ", "ФЗ № 223-ФЗ О закупках товаров, работ, услуг отдельными видами юридических лиц", ["Федеральный закон № 223-ФЗ О закупках товаров, работ, услуг отдельными видами юридических лиц"]),
    ("115-ФЗ", "ФЗ № 115-ФЗ О противодействии легализации (отмыванию) доходов", ["Федеральный закон № 115-ФЗ О противодействии легализации (отмыванию) доходов, полученных преступным путем, и финансированию терроризма"]),
    ("102-ФЗ", "ФЗ № 102-ФЗ Об ипотеке (залоге недвижимости)", ["Федеральный закон № 102-ФЗ Об ипотеке (залоге недвижимости)"]),
    ("422-ФЗ", "ФЗ № 422-ФЗ О специальном налоговом режиме для самозанятых", ["Федеральный закон № 422-ФЗ О проведении эксперимента по установлению специального налогового режима Налог на профессиональный доход"]),
    ("149-ФЗ", "ФЗ № 149-ФЗ Об информации, информационных технологиях и о защите информации", ["Федеральный закон № 149-ФЗ Об информации, информационных технологиях и о защите информации"]),
    ("248-ФЗ", "ФЗ № 248-ФЗ О государственном контроле (надзоре) и муниципальном контроле в РФ", ["Федеральный закон № 248-ФЗ О государственном контроле (надзоре) и муниципальном контроле в Российской Федерации"]),
    ("161-ФЗ", "ФЗ № 161-ФЗ О национальной платёжной системе", ["Федеральный закон № 161-ФЗ О национальной платежной системе"]),
    ("7-ФЗ", "ФЗ № 7-ФЗ Об охране окружающей среды", ["Федеральный закон № 7-ФЗ Об охране окружающей среды"]),
    ("89-ФЗ", "ФЗ № 89-ФЗ Об отходах производства и потребления", ["Федеральный закон № 89-ФЗ Об отходах производства и потребления"]),
    ("52-ФЗ", "ФЗ № 52-ФЗ О санитарно-эпидемиологическом благополучии населения", ["Федеральный закон № 52-ФЗ О санитарно-эпидемиологическом благополучии населения"]),
    ("69-ФЗ", "ФЗ № 69-ФЗ О пожарной безопасности", ["Федеральный закон № 69-ФЗ О пожарной безопасности"]),
    ("384-ФЗ", "ФЗ № 384-ФЗ Технический регламент о безопасности зданий и сооружений", ["Федеральный закон № 384-ФЗ Технический регламент о безопасности зданий и сооружений"]),
    ("259-ФЗ", "ФЗ № 259-ФЗ Устав автомобильного транспорта и городского наземного электрического транспорта", ["Федеральный закон № 259-ФЗ Устав автомобильного транспорта и городского наземного электрического транспорта"]),
    ("196-ФЗ", "ФЗ № 196-ФЗ О безопасности дорожного движения", ["Федеральный закон № 196-ФЗ О безопасности дорожного движения"]),
    ("16-ФЗ", "ФЗ № 16-ФЗ О транспортной безопасности", ["Федеральный закон № 16-ФЗ О транспортной безопасности"]),
]


@app.get("/api/admin/all")
def admin_all(user=Depends(_auth)):
    _admin(user)
    conn = get_connection()

    def safe(fn, default):
        try:
            return fn()
        except Exception:
            return default

    cols = [r["name"] for r in conn.execute("PRAGMA table_info(users)").fetchall()]
    sel = "id, email, tariff, checks_left, verified" + (", created_at" if "created_at" in cols else "")
    out = {}
    out["users"] = safe(lambda: [dict(r) for r in conn.execute(
        f"SELECT {sel} FROM users ORDER BY id DESC")], [])
    out["stats"] = safe(lambda: {
        "users": conn.execute("SELECT COUNT(*) FROM users").fetchone()[0],
        "verified": conn.execute("SELECT COUNT(*) FROM users WHERE verified=1").fetchone()[0],
        "analyses": conn.execute("SELECT COUNT(*) FROM analyses").fetchone()[0],
        "contracts": conn.execute("SELECT COUNT(*) FROM contracts").fetchone()[0],
        "laws": conn.execute("SELECT COUNT(*) FROM laws").fetchone()[0]}, {})
    out["stats"]["feedback_avg"] = safe(lambda: conn.execute(
        "SELECT ROUND(AVG(rating),2) FROM feedback").fetchone()[0], None)
    out["stats"]["support_open"] = safe(lambda: conn.execute(
        "SELECT COUNT(*) FROM support_messages WHERE replied=0").fetchone()[0], 0)
    out["regs"] = safe(lambda: [dict(r) for r in conn.execute(
        "SELECT date(created_at) d, COUNT(*) c FROM users GROUP BY d ORDER BY d")], [])
    out["anls"] = safe(lambda: [dict(r) for r in conn.execute(
        "SELECT date(created_at) d, COUNT(*) c FROM analyses GROUP BY d ORDER BY d")], [])
    out["paid"] = safe(lambda: conn.execute(
        "SELECT COUNT(*) FROM users WHERE tariff!='Free'").fetchone()[0], 0)
    out["promos"] = safe(lambda: [dict(r) for r in conn.execute(
        "SELECT * FROM promocodes ORDER BY id DESC")], [])
    out["top_promos"] = sorted(out["promos"], key=lambda p: p.get("used_count") or 0, reverse=True)[:5]
    out["support"] = safe(lambda: [dict(r) for r in conn.execute(
        "SELECT * FROM support_messages ORDER BY id DESC LIMIT 50")], [])
    out["feedback"] = safe(lambda: [dict(r) for r in conn.execute(
        "SELECT f.*, u.email FROM feedback f LEFT JOIN users u ON u.id=f.user_id "
        "ORDER BY f.id DESC LIMIT 100")], [])
    out["consults"] = safe(lambda: [dict(r) for r in conn.execute(
        "SELECT c.*, u.email FROM consult_requests c LEFT JOIN users u ON u.id=c.user_id "
        "ORDER BY c.id DESC LIMIT 100")], [])
    out["laws_list"] = safe(lambda: [dict(r) for r in conn.execute(
        "SELECT code, title, LENGTH(COALESCE(full_text,'')) ft FROM laws ORDER BY code")], [])
    out["job"] = dict(globals().get("_LAWS_JOB", {}))
    conn.close()
    return out


class LawsReloadIn(BaseModel):
    extended: bool = False


@app.post("/api/admin/laws_reload")
def admin_laws_reload2(data: LawsReloadIn, user=Depends(_auth)):
    _admin(user)
    import threading
    job = globals().setdefault("_LAWS_JOB", {"running": False, "done": 0, "total": 0, "last": "", "error": ""})
    if job["running"]:
        return {"ok": False, "detail": "Уже выполняется"}

    def run():
        try:
            from core.laws_autoload import DEFAULT_PACK, autoload_law
            pack = list(DEFAULT_PACK) + (list(EXTENDED_PACK) if data.extended else [])
            job.update(running=True, done=0, total=len(pack), last="", error="")
            for prefix, title, cands in pack:
                n, err, source = autoload_law(prefix, title, cands)
                job["done"] += 1
                job["last"] = f"{title}: статей {n}" if n else f"{title}: {err}"
        except Exception as e:
            job["error"] = str(e)
        finally:
            job["running"] = False

    threading.Thread(target=run, daemon=True).start()
    return {"ok": True}


AUDIO_TIERS = {"Standard", "Pro", "Business", "Business Pro"}


@app.get("/api/analyses/{aid}/audio")
def analysis_audio(aid: int, user=Depends(_auth)):
    if user["tariff"] not in AUDIO_TIERS:
        raise HTTPException(403, "Аудиоверсия доступна на тарифах Standard, Pro, Business и Business Pro")
    import os
    cache_dir = "/opt/app/cache/audio"
    os.makedirs(cache_dir, exist_ok=True)
    path = os.path.join(cache_dir, f"{aid}.mp3")
    if os.path.exists(path):
        with open(path, "rb") as f:
            audio = f.read()
        return Response(content=audio, media_type="audio/mpeg")
    conn = get_connection()
    row = conn.execute("SELECT report FROM analyses WHERE id = ? AND user_id = ?",
                       (aid, user["id"])).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "Отчёт не найден")
    text = (row["report"] or "").split("HIGHLIGHTS_JSON:")[0][:6000]
    audio = None
    try:
        from gtts import gTTS
        import io
        buf = io.BytesIO()
        gTTS(text=text, lang="ru").write_to_fp(buf)
        audio = buf.getvalue()
    except Exception:
        audio = None
    if not audio:
        raise HTTPException(500, "Генерация аудио временно недоступна")
    with open(path, "wb") as f:
        f.write(audio)
    return Response(content=audio, media_type="audio/mpeg")


COMPARE_TIERS = {"Pro", "Business", "Business Pro"}
COMPARE_KEYS = ["неустойк", "штраф", "пеня", "размер", "цена", "срок", "уведомл",
                "ответствен", "залог", "обеспеч", "расторж", "отказ", "продлен",
                "индексац", "суд", "арбитраж", "подсудност", "арендн", "оплат"]


class CompareIn(BaseModel):
    old_text: str
    new_text: str


@app.post("/api/compare")
def compare_versions(data: CompareIn, user=Depends(_auth)):
    if user["tariff"] not in COMPARE_TIERS:
        raise HTTPException(403, "Сравнение версий доступно на тарифах Pro, Business и Business Pro")
    if user["checks_left"] < 1:
        raise HTTPException(402, "Недостаточно проверок")
    if len(data.old_text) > 200000 or len(data.new_text) > 200000:
        raise HTTPException(413, "Каждая версия должна быть не длиннее 200 000 символов")
    import difflib
    import re

    def clauses(t):
        parts = re.split(r"\n(?=\s*\d+[.)]\s)", t or "")
        return [p.strip() for p in parts if p.strip()]

    old_c = clauses(data.old_text)
    new_c = clauses(data.new_text)
    sm = difflib.SequenceMatcher(None, old_c, new_c)
    changes = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            continue
        if tag == "delete":
            changes += [{"type": "removed", "old": old_c[k][:1500], "new": ""} for k in range(i1, i2)]
        elif tag == "insert":
            changes += [{"type": "added", "old": "", "new": new_c[k][:1500]} for k in range(j1, j2)]
        else:
            for n, k in enumerate(range(i1, i2)):
                changes.append({"type": "changed", "old": old_c[k][:1500],
                                "new": (new_c[j1 + n][:1500] if j1 + n < j2 else "")})
    for c in changes:
        blob = (c["old"] + " " + c["new"]).lower()
        c["important"] = any(k in blob for k in COMPARE_KEYS)
    conn = get_connection()
    conn.execute("UPDATE users SET checks_left = MAX(0, checks_left - 1) WHERE id = ?",
                 (user["id"],))
    conn.commit()
    conn.close()
    stats = {"added": sum(1 for c in changes if c["type"] == "added"),
             "removed": sum(1 for c in changes if c["type"] == "removed"),
             "changed": sum(1 for c in changes if c["type"] == "changed"),
             "important": sum(1 for c in changes if c["important"])}
    changes.sort(key=lambda c: (not c["important"], c["type"]))
    return {"stats": stats, "changes": changes[:300]}
    

@app.post("/api/delete_account")
def delete_account(user=Depends(_auth)):
    from utils.auth import delete_user
    delete_user(user["id"])
    return {"ok": True}


class RedlineV2In(BaseModel):
    analysis_id: int
    scenario: str = ""


@app.post("/api/tools/redline_v2")
def tool_redline_v2(data: RedlineV2In, user=Depends(_auth)):
    conn = get_connection()
    row = conn.execute(
        "SELECT a.highlights FROM analyses a WHERE a.id = ? AND a.user_id = ?",
        (data.analysis_id, user["id"])).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "Отчёт не найден")
    import json as _json
    try:
        hl = _json.loads(row["highlights"] or "[]")
    except Exception:
        hl = []
    risks = [x for x in hl if x.get("level") in ("red", "yellow")][:12]
    if not risks:
        return {"text": "Существенных рисков не найдено — правки не требуются."}
    from core.analyzer import ask_deepseek
    prompt = (
        "Ты — юрист, готовящий ПРОТОКОЛ ИЗМЕНЕНИЙ к договору для переговоров. "
        "По каждому риску ниже выдай ровно 4 строки:\n"
        "ПУНКТ: <номер пункта договора>\n"
        "СТАРЫЙ: <короткая цитата вредной формулировки>\n"
        "НОВЫЙ: <полная безопасная замена пункта, юридическим стилем>\n"
        "ОСНОВАНИЕ: <статья ГК/иного закона или «прямой нормы в базе нет»>\n"
        "Риски:\n" +
        "\n".join(f"- {x.get('quote','')[:200]} | {x.get('reason','')[:200]}" for x in risks) +
        "\nВыведи только протокол, без вступления и заключения. Язык русский."
    )
    text = ask_deepseek(prompt, max_tokens=3000, temperature=0.2)
    return {"text": text}


class FeedbackV2In(BaseModel):
    analysis_id: int
    rating: int
    comment: str = ""


@app.post("/api/feedback_v2")
def feedback_v2(data: FeedbackV2In, user=Depends(_auth)):
    conn = get_connection()
    conn.execute("DELETE FROM feedback WHERE analysis_id = ? AND user_id = ?",
                 (data.analysis_id, user["id"]))
    conn.execute("INSERT INTO feedback (analysis_id, user_id, rating, comment) VALUES (?,?,?,?)",
                 (data.analysis_id, user["id"], max(1, min(5, data.rating)), data.comment))
    conn.commit()
    conn.close()
    return {"ok": True}


def _lawyer_used(email):
    from datetime import date
    conn = get_connection()
    conn.execute("CREATE TABLE IF NOT EXISTS lawyer_usage (email TEXT, day TEXT, used INTEGER, PRIMARY KEY (email, day))")
    row = conn.execute("SELECT used FROM lawyer_usage WHERE email = ? AND day = ?",
                       (email, date.today().isoformat())).fetchone()
    conn.close()
    return row["used"] if row else 0


def _lawyer_inc(email):
    from datetime import date
    conn = get_connection()
    conn.execute("CREATE TABLE IF NOT EXISTS lawyer_usage (email TEXT, day TEXT, used INTEGER, PRIMARY KEY (email, day))")
    conn.execute("INSERT INTO lawyer_usage (email, day, used) VALUES (?, ?, 1) "
                 "ON CONFLICT(email, day) DO UPDATE SET used = used + 1",
                 (email, date.today().isoformat()))
    conn.commit()
    conn.close()


@app.get("/api/lawyer_left")
def lawyer_left(user=Depends(_auth)):
    limit = LAWYER_LIMITS.get(user["tariff"], 0)
    left = max(0, limit - _lawyer_used(user["email"]))
    return {"limit": limit, "left": left}


ANALYSIS_LAWYER_LIMITS = {"Standard": 10, "Pro": 20, "Business": 15, "Business Pro": 35}


@app.get("/api/lawyer_chats")
def lawyer_chats(user=Depends(_auth)):
    conn = get_connection()
    conn.execute("CREATE TABLE IF NOT EXISTS lawyer_chats (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, analysis_id INTEGER, question TEXT, answer TEXT, created_at TEXT DEFAULT (datetime('now')))")
    rows = [dict(r) for r in conn.execute(
        "SELECT id, analysis_id, question, answer, created_at FROM lawyer_chats WHERE user_id = ? ORDER BY id DESC LIMIT 50",
        (user["id"],))]
    conn.close()
    return rows


@app.get("/api/lawyer_analysis_left")
def lawyer_analysis_left(analysis_id: int, user=Depends(_auth)):
    limit = ANALYSIS_LAWYER_LIMITS.get(user["tariff"], 0)
    conn = get_connection()
    conn.execute("CREATE TABLE IF NOT EXISTS lawyer_analysis_usage (analysis_id INTEGER PRIMARY KEY, used INTEGER DEFAULT 0)")
    row = conn.execute("SELECT used FROM lawyer_analysis_usage WHERE analysis_id = ?", (analysis_id,)).fetchone()
    conn.close()
    used = row["used"] if row else 0
    return {"limit": limit, "left": max(0, limit - used)}


class LawyerAnalysisIn(BaseModel):
    analysis_id: int
    question: str


@app.post("/api/lawyer_analysis")
def lawyer_analysis(data: LawyerAnalysisIn, user=Depends(_auth)):
    limit = ANALYSIS_LAWYER_LIMITS.get(user["tariff"], 0)
    if limit == 0:
        raise HTTPException(403, "Вопросы юристу по анализу доступны с тарифа Standard")
    conn = get_connection()
    conn.execute("CREATE TABLE IF NOT EXISTS lawyer_analysis_usage (analysis_id INTEGER PRIMARY KEY, used INTEGER DEFAULT 0)")
    row = conn.execute("SELECT used FROM lawyer_analysis_usage WHERE analysis_id = ?", (data.analysis_id,)).fetchone()
    used = row["used"] if row else 0
    if used >= limit:
        conn.close()
        raise HTTPException(429, f"Лимит вопросов юристу по этому договору исчерпан: {limit} на тарифе {user['tariff']}")
    a = conn.execute(
        "SELECT a.report, c.text FROM analyses a JOIN contracts c ON c.id = a.contract_id "
        "WHERE a.id = ? AND a.user_id = ?", (data.analysis_id, user["id"])).fetchone()
    if not a:
        conn.close()
        raise HTTPException(404, "Отчёт не найден")
    conn.execute("INSERT INTO lawyer_analysis_usage (analysis_id, used) VALUES (?,1) "
                 "ON CONFLICT(analysis_id) DO UPDATE SET used = used + 1", (data.analysis_id,))
    conn.commit()
    conn.close()
    from core.analyzer import lawyer247_stream
    context = ("КОНТЕКСТ: текст договора (фрагмент):\n" + (a["text"] or "")[:6000] +
               "\n\nФРАГМЕНТ ГОТОВОГО ОТЧЁТА ПО ЭТОМУ ДОГОВОРУ:\n" +
               (a["report"] or "").split("HIGHLIGHTS_JSON:")[0][:4000] +
               "\n\nВОПРОС ПО ЭТОМУ ДОГОВОРУ: ")
    gen, model = lawyer247_stream(context + "Отвечай строго на русском языке, без иноязычных вставок. " + data.question, [], user["tariff"])

    def stream():
        out = []
        for ch in gen:
            out.append(ch)
            yield f"data: {json.dumps({'chunk': ch}, ensure_ascii=False)}\n\n"
        try:
            conn2 = get_connection()
            conn2.execute("CREATE TABLE IF NOT EXISTS lawyer_chats (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, analysis_id INTEGER, question TEXT, answer TEXT, created_at TEXT DEFAULT (datetime('now')))")
            conn2.execute("INSERT INTO lawyer_chats (user_id, analysis_id, question, answer) VALUES (?,?,?,?)",
                          (user["id"], data.analysis_id, data.question, "".join(out)))
            conn2.commit()
            conn2.close()
        except Exception:
            pass
        yield f"data: {json.dumps({'done': True, 'left': limit - used - 1}, ensure_ascii=False)}\n\n"

    return StreamingResponse(stream(), media_type="text/event-stream")


@app.delete("/api/analyses/{aid}")
def analysis_delete(aid: int, user=Depends(_auth)):
    conn = get_connection()
    row = conn.execute("SELECT contract_id FROM analyses WHERE id = ? AND user_id = ?",
                       (aid, user["id"])).fetchone()
    if not row:
        conn.close()
        raise HTTPException(404, "Отчёт не найден")
    conn.execute("DELETE FROM analyses WHERE id = ? AND user_id = ?", (aid, user["id"]))
    try:
        conn.execute("DELETE FROM lawyer_chats WHERE analysis_id = ?", (aid,))
        conn.execute("DELETE FROM lawyer_analysis_usage WHERE analysis_id = ?", (aid,))
        conn.execute("DELETE FROM contracts WHERE id = ? AND NOT EXISTS "
                     "(SELECT 1 FROM analyses a WHERE a.contract_id = ?)",
                     (row["contract_id"], row["contract_id"]))
    except Exception:
        pass
    conn.commit()
    conn.close()
    return {"ok": True}


@app.delete("/api/lawyer_chats/{cid}")
def lawyer_chat_delete(cid: int, user=Depends(_auth)):
    conn = get_connection()
    conn.execute("DELETE FROM lawyer_chats WHERE id = ? AND user_id = ?", (cid, user["id"]))
    conn.commit()
    conn.close()
    return {"ok": True}


@app.get("/api/admin/laws_search")
def admin_laws_search(q: str = "", user=Depends(_auth)):
    _admin(user)
    q = (q or "").strip()
    if not q:
        return []
    import re as _re
    conn = get_connection()
    rows = conn.execute("SELECT code, title, essence, full_text FROM laws").fetchall()
    conn.close()
    out = []
    m = _re.match(r"^\s*(?P<code>\d+\s*-\s*ФЗ|[A-Za-zА-Яа-я][A-Za-zА-Яа-я\d\-]*)?\s*"
                  r"(?:ст\.?|статья)?\s*(?P<num>\d+(?:[.\d]+)?)\s*$", q, _re.I)
    if m and m.group("num"):
        code_q = (m.group("code") or "").upper().replace(" ", "")
        num = m.group("num")
        for r in rows:
            c = (r["code"] or "").upper()
            base, _, art = c.rsplit(" ", 1) if " " in c else (c, "", "")
            if art == num and (not code_q or code_q in base):
                out.append({"code": r["code"], "art": num, "loaded": True,
                            "snippet": " ".join(((r["title"] or "") + " " +
                                                 (r["full_text"] or r["essence"] or ""))[:400].split())})
                if len(out) >= 12:
                    break
        if not out:
            out.append({"code": (m.group("code") or "ВСЕ КОДЕКСЫ").upper(), "art": num,
                        "loaded": False,
                        "snippet": "Статья не найдена ни в одном загруженном источнике"})
        return out
    low = q.lower()
    for r in rows:
        blob = ((r["code"] or "") + " " + (r["title"] or "") + " " + (r["full_text"] or "")).lower()
        pos = blob.find(low)
        if pos >= 0 and len(out) < 20:
            out.append({"code": r["code"],
                        "art": (r["code"] or "").rsplit(" ", 1)[-1], "loaded": True,
                        "snippet": " ".join(blob[max(0, pos - 60):pos + 340].split())})
    return out
import threading

_LAWS_JOB_V2 = {"running": False, "done": 0, "total": 0, "current": "", "log": []}

ADD_PACK_V2 = [
    ("КОНСТ", "Конституция Российской Федерации", ["Конституция Российской Федерации"]),
    ("54-ФЗ", "ФЗ № 54-ФЗ О применении контрольно-кассовой техники",
     ["Федеральный закон № 54-ФЗ О применении контрольно-кассовой техники при осуществлении расчетов в Российской Федерации"]),
]


def _fetch_wikisource_v2(title):
    import urllib.request
    import urllib.parse
    import urllib.error
    import json as _j
    import time as _t
    url = ("https://ru.wikisource.org/w/api.php?action=parse&redirects=1&page="
           + urllib.parse.quote(title) + "&prop=wikitext&format=json")
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                             "(KHTML, like Gecko) Chrome/124.0 Safari/537.36 MCA-LawsBot/1.0",
               "Accept": "application/json"}
    for attempt in range(4):
        try:
            req = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(req, timeout=60) as r:
                d = _j.load(r)
            txt = d.get("parse", {}).get("wikitext", {}).get("*", "") or ""
            _t.sleep(1.2)
            return txt
        except urllib.error.HTTPError as e:
            wait = int(e.headers.get("Retry-After") or 0) or (5 * (attempt + 1))
            _t.sleep(min(wait, 30))
        except Exception:
            _t.sleep(3 * (attempt + 1))
    return ""


def _search_title_v2(query):
    import urllib.request
    import urllib.parse
    import urllib.error
    import json as _j
    import time as _t
    url = ("https://ru.wikisource.org/w/api.php?action=query&list=search&srsearch="
           + urllib.parse.quote(query or "") + "&srlimit=1&format=json")
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                             "(KHTML, like Gecko) Chrome/124.0 Safari/537.36 MCA-LawsBot/1.0"}
    for attempt in range(3):
        try:
            req = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(req, timeout=30) as r:
                d = _j.load(r)
            res = d.get("query", {}).get("search", [])
            _t.sleep(1.0)
            return res[0]["title"] if res else ""
        except Exception:
            _t.sleep(3 * (attempt + 1))
    return ""


def _fetch_law_full_v2(title):
    import re as _re
    ft = _fetch_wikisource_v2(title)
    if len(ft) < 2000:
        return ""
    if _re.search(r"Статья\s*\d", ft):
        return ft
    links = _re.findall(r"\[\[" + _re.escape(title) + r"/([^\]|#]+)(?:\|[^\]]*)?\]\]", ft)
    if not links:
        full = _re.findall(r"\[\[(" + _re.escape(title) + r"/[^\]|#]+)(?:\|[^\]]*)?\]\]", ft)
        links = [x[len(title) + 1:] for x in full]
    seen, parts = set(), [ft]
    for sub in links:
        sub = sub.strip()
        if not sub or sub in seen:
            continue
        seen.add(sub)
        chunk = _fetch_wikisource_v2(title + "/" + sub)
        if len(chunk) > 500:
            parts.append(chunk)
        if len(seen) >= 80:
            break
    return "\n".join(parts) if len(parts) > 1 else ft


def _split_articles_v2(ft, code, title):
    import re as _re
    ft = _re.sub(r"\{\{[^}]*\}\}", "", ft or "")
    ft = _re.sub(r"<!--.*?-->", "", ft, flags=_re.S)
    ms = list(_re.finditer(r"Статья\s*(\d+(?:[.\d]+)?)", ft))
    rows = []
    for i, m in enumerate(ms):
        start = m.start()
        end = ms[i + 1].start() if i + 1 < len(ms) else min(len(ft), start + 8000)
        chunk = ft[start:end].strip()
        if len(chunk) < 40:
            continue
        rows.append((f"{code} {m.group(1)}", f"{title} — ст. {m.group(1)}", chunk[:6000]))
    return rows


def _run_pack_v2(items):
    import re as _re
    J = _LAWS_JOB_V2
    J.update(running=True, done=0, total=len(items), current="", log=[])
    conn = get_connection()
    conn.execute("CREATE TABLE IF NOT EXISTS laws (id INTEGER PRIMARY KEY AUTOINCREMENT, "
                 "code TEXT UNIQUE, title TEXT, essence TEXT, tags TEXT, category TEXT, "
                 "full_text TEXT DEFAULT '')")
    for code, title, cands in items:
        J["current"] = title
        try:
            ft = ""
            for c in list(cands) + [title]:
                ft = _fetch_law_full_v2(c)
                if ft and _re.search(r"Статья\s*\d", ft):
                    break
                if not ft:
                    resolved = _search_title_v2(c)
                    if resolved:
                        ft = _fetch_law_full_v2(resolved)
                        if ft and _re.search(r"Статья\s*\d", ft):
                            break
            if not ft:
                J["log"].append(f"{code}: не удалось скачать")
                J["done"] += 1
                continue
            rows = _split_articles_v2(ft, code, title)
            if not rows:
                J["log"].append(f"{code}: текст есть, но статей не найдено")
                J["done"] += 1
                continue
            tag = _re.split(r"[\s-]", code)[0].lower()
            for c2, t2, chunk in rows:
                conn.execute("INSERT OR REPLACE INTO laws "
                             "(code, title, essence, tags, category, full_text) "
                             "VALUES (?,?,?,?,?,?)",
                             (c2, t2, "", tag + " закон", "закон", chunk))
            conn.commit()
            J["log"].append(f"{code}: статей {len(rows)}")
        except Exception as e:
            J["log"].append(f"{code}: ошибка {e}")
        J["done"] += 1
    conn.close()
    J["running"] = False
    J["current"] = ""
    ("ГК", "Гражданский кодекс РФ (часть вторая)", ["Гражданский кодекс Российской Федерации (часть вторая)"]),
    ("ТК", "Трудовой кодекс РФ", ["Трудовой кодекс Российской Федерации"]),
    ("ЗоЗПП", "Закон О защите прав потребителей", ["Закон РФ О защите прав потребителей"]),
    ("КоАП", "Кодекс об административных правонарушениях РФ", ["Кодекс Российской Федерации об административных правонарушениях"]),
    ("ГПК", "Гражданский процессуальный кодекс РФ", ["Гражданский процессуальный кодекс Российской Федерации"]),
    ("АПК", "Арбитражный процессуальный кодекс РФ", ["Арбитражный процессуальный кодекс Российской Федерации"]),
    ("СК", "Семейный кодекс РФ", ["Семейный кодекс Российской Федерации"]),
    ("152-ФЗ", "ФЗ О персональных данных", ["Федеральный закон О персональных данных"]),
    ("40-ФЗ", "ФЗ Об ОСАГО", ["Федеральный закон Об обязательном страховании гражданской ответственности владельцев транспортных средств"]),

EXT_PACK_V2 = [
    ("ЖК", "Жилищный кодекс РФ", ["Жилищный кодекс Российской Федерации"]),
    ("ЗК", "Земельный кодекс РФ", ["Земельный кодекс Российской Федерации"]),
    ("НК-1", "Налоговый кодекс РФ (часть 1)", ["Налоговый кодекс Российской Федерации (часть первая)"]),
    ("НК-2", "Налоговый кодекс РФ (часть 2)", ["Налоговый кодекс Российской Федерации (часть вторая)"]),
    ("ГК-1", "Гражданский кодекс РФ (часть первая)", ["Гражданский кодекс Российской Федерации (часть первая)"]),
    ("ГК-3", "Гражданский кодекс РФ (часть третья)", ["Гражданский кодекс Российской Федерации (часть третья)"]),
    ("ГК-4", "Гражданский кодекс РФ (часть четвёртая)", ["Гражданский кодекс Российской Федерации (часть четвертая)"]),
    ("УК", "Уголовный кодекс РФ", ["Уголовный кодекс Российской Федерации"]),
    ("ВК", "Водный кодекс РФ", ["Водный кодекс Российской Федерации"]),
    ("ЛК", "Лесной кодекс РФ", ["Лесной кодекс Российской Федерации"]),
    ("КТМ", "Кодекс торгового мореплавания РФ", ["Кодекс торгового мореплавания Российской Федерации"]),
    ("КВВТ", "Кодекс внутреннего водного транспорта РФ", ["Кодекс внутреннего водного транспорта Российской Федерации"]),
    ("УЖТ", "Устав железнодорожного транспорта РФ", ["Устав железнодорожного транспорта Российской Федерации"]),
    ("353-ФЗ", "ФЗ № 353-ФЗ О потребительском кредите (займе)", ["Федеральный закон № 353-ФЗ О потребительском кредите (займе)"]),
    ("214-ФЗ", "ФЗ № 214-ФЗ Об участии в долевом строительстве", ["Федеральный закон № 214-ФЗ Об участии в долевом строительстве многоквартирных домов и иных объектов недвижимости"]),
    ("127-ФЗ", "ФЗ № 127-ФЗ О несостоятельности (банкротстве)", ["Федеральный закон № 127-ФЗ О несостоятельности (банкротстве)"]),
    ("63-ФЗ", "ФЗ № 63-ФЗ Об электронной подписи", ["Федеральный закон № 63-ФЗ Об электронной подписи"]),
    ("98-ФЗ", "ФЗ № 98-ФЗ О коммерческой тайне", ["Федеральный закон № 98-ФЗ О коммерческой тайне"]),
    ("38-ФЗ", "ФЗ № 38-ФЗ О рекламе", ["Федеральный закон № 38-ФЗ О рекламе"]),
    ("164-ФЗ", "ФЗ № 164-ФЗ О финансовой аренде (лизинге)", ["Федеральный закон № 164-ФЗ О финансовой аренде (лизинге)"]),
    ("229-ФЗ", "ФЗ № 229-ФЗ Об исполнительном производстве", ["Федеральный закон № 229-ФЗ Об исполнительном производстве"]),
    ("135-ФЗ", "ФЗ № 135-ФЗ О защите конкуренции", ["Федеральный закон № 135-ФЗ О защите конкуренции"]),
    ("44-ФЗ", "ФЗ № 44-ФЗ О контрактной системе в сфере закупок", ["Федеральный закон № 44-ФЗ О контрактной системе в сфере закупок товаров, работ, услуг для обеспечения государственных и муниципальных нужд"]),
    ("223-ФЗ", "ФЗ № 223-ФЗ О закупках отдельными видами юрлиц", ["Федеральный закон № 223-ФЗ О закупках товаров, работ, услуг отдельными видами юридических лиц"]),
    ("115-ФЗ", "ФЗ № 115-ФЗ О противодействии легализации доходов", ["Федеральный закон № 115-ФЗ О противодействии легализации (отмыванию) доходов, полученных преступным путем, и финансированию терроризма"]),
    ("102-ФЗ", "ФЗ № 102-ФЗ Об ипотеке (залоге недвижимости)", ["Федеральный закон № 102-ФЗ Об ипотеке (залоге недвижимости)"]),
    ("422-ФЗ", "ФЗ № 422-ФЗ О налоге на профессиональный доход", ["Федеральный закон № 422-ФЗ О проведении эксперимента по установлению специального налогового режима Налог на профессиональный доход"]),
    ("149-ФЗ", "ФЗ № 149-ФЗ Об информации и информзащите", ["Федеральный закон № 149-ФЗ Об информации, информационных технологиях и о защите информации"]),
    ("248-ФЗ", "ФЗ № 248-ФЗ О государственном контроле (надзоре)", ["Федеральный закон № 248-ФЗ О государственном контроле (надзоре) и муниципальном контроле в Российской Федерации"]),
    ("161-ФЗ", "ФЗ № 161-ФЗ О национальной платёжной системе", ["Федеральный закон № 161-ФЗ О национальной платежной системе"]),
    ("7-ФЗ", "ФЗ № 7-ФЗ Об охране окружающей среды", ["Федеральный закон № 7-ФЗ Об охране окружающей среды"]),
    ("89-ФЗ", "ФЗ № 89-ФЗ Об отходах производства и потребления", ["Федеральный закон № 89-ФЗ Об отходах производства и потребления"]),
    ("52-ФЗ", "ФЗ № 52-ФЗ О санитарно-эпидемиологическом благополучии", ["Федеральный закон № 52-ФЗ О санитарно-эпидемиологическом благополучии населения"]),
    ("69-ФЗ", "ФЗ № 69-ФЗ О пожарной безопасности", ["Федеральный закон № 69-ФЗ О пожарной безопасности"]),
    ("384-ФЗ", "ФЗ № 384-ФЗ Техрегламент о безопасности зданий", ["Федеральный закон № 384-ФЗ Технический регламент о безопасности зданий и сооружений"]),
    ("259-ФЗ", "ФЗ № 259-ФЗ Устав автомобильного транспорта", ["Федеральный закон № 259-ФЗ Устав автомобильного транспорта и городского наземного электрического транспорта"]),
    ("196-ФЗ", "ФЗ № 196-ФЗ О безопасности дорожного движения", ["Федеральный закон № 196-ФЗ О безопасности дорожного движения"]),
    ("16-ФЗ", "ФЗ № 16-ФЗ О транспортной безопасности", ["Федеральный закон № 16-ФЗ О транспортной безопасности"]),
]
class LawsReloadV2In(BaseModel):
    pack: str = "base"


@app.post("/api/admin/laws_reload_v2")
def laws_reload_v2(data: LawsReloadV2In, user=Depends(_auth)):
    _admin(user)
    if _LAWS_JOB_V2["running"]:
        return {"started": False, "reason": "already running"}
    try:
        from core.law_packs import BASE_PACK_V2 as _B, EXT_PACK_V2 as _E, ADD_PACK_V2 as _A
    except Exception:
        _B, _E, _A = [], [], []
    _B = _B or globals().get("BASE_PACK_V2") or []
    _E = _E or globals().get("EXT_PACK_V2") or []
    _A = _A or globals().get("ADD_PACK_V2") or []
    items = list(_B if data.pack == "base" else _E + _A)
    if not items:
        raise HTTPException(500, "Паки не найдены: создай core/law_packs.py")
    threading.Thread(target=_run_pack_v2, args=(items,), daemon=True).start()
    return {"started": True, "total": len(items)}


@app.get("/api/admin/laws_job_v2")
def laws_job_v2(user=Depends(_auth)):
    _admin(user)
    J = _LAWS_JOB_V2
    return {"running": J["running"], "done": J["done"], "total": J["total"],
            "current": J["current"], "log": J["log"][-10:],
            "last": (not J["running"]) and J["done"] > 0}


@app.get("/api/admin/laws_search_v2")
def laws_search_v2(q: str = "", user=Depends(_auth)):
    _admin(user)
    q = (q or "").strip()
    if not q:
        return []
    import re as _re
    conn = get_connection()
    rows = conn.execute("SELECT code, title, essence, full_text FROM laws").fetchall()
    conn.close()
    out = []
    m = _re.match(r"^\s*(?P<code>\d+\s*-\s*ФЗ|[A-Za-zА-Яа-я][A-Za-zА-Яа-я\d\-]*)?\s*"
                  r"(?:ст\.?|статья)?\s*(?P<num>\d+(?:[.\d]+)?)\s*$", q, _re.I)
    if m and m.group("num"):
        code_q = (m.group("code") or "").upper().replace(" ", "")
        num = m.group("num")
        for r in rows:
            c = (r["code"] or "").upper()
            base, _, art = c.rsplit(" ", 1) if " " in c else (c, "", "")
            if art == num and (not code_q or code_q in base):
                out.append({"code": r["code"], "art": num, "loaded": True,
                            "snippet": " ".join(((r["title"] or "") + " " +
                                                 (r["full_text"] or r["essence"] or ""))[:400].split())})
                if len(out) >= 12:
                    break
        if not out:
            out.append({"code": (m.group("code") or "ВСЕ КОДЕКСЫ").upper(), "art": num,
                        "loaded": False,
                        "snippet": "Статья не найдена ни в одном загруженном источнике"})
        return out
    low = q.lower()
    for r in rows:
        blob = ((r["code"] or "") + " " + (r["title"] or "") + " " + (r["full_text"] or "")).lower()
        pos = blob.find(low)
        if pos >= 0 and len(out) < 20:
            out.append({"code": r["code"],
                        "art": (r["code"] or "").rsplit(" ", 1)[-1], "loaded": True,
                        "snippet": " ".join(blob[max(0, pos - 60):pos + 340].split())})
    return outclass AdminTariffIn(BaseModel):
    user_id: int
    tariff: str
    promo: str = ""


class AdminChecksIn(BaseModel):
    user_id: int
    amount: int


class AdminPromoIn(BaseModel):
    code: str
    percent: int
    max_uses: int = 100


class AdminPromoToggleIn(BaseModel):
    code: str


class AdminSupportReplyIn(BaseModel):
    message_id: int
    text: str


class AdminConsultReplyIn(BaseModel):
    id: int
    text: str


class AdminUserIdIn(BaseModel):
    user_id: int


TIER_CHECKS_ADMIN = {"Free": 3, "Light": 35, "Standard": 65, "Pro": 125,
                     "Business": 150, "Business Pro": 200}


@app.post("/api/admin/tariff")
def admin_tariff(data: AdminTariffIn, user=Depends(_auth)):
    _admin(user)
    add = TIER_CHECKS_ADMIN.get(data.tariff, 0)
    conn = get_connection()
    conn.execute("UPDATE users SET tariff = ?, checks_left = COALESCE(checks_left, 0) + ? WHERE id = ?",
                 (data.tariff, add, data.user_id))
    if data.promo:
        try:
            conn.execute("UPDATE promos SET used_count = COALESCE(used_count, 0) + 1 WHERE code = ?",
                         (data.promo.upper(),))
        except Exception:
            pass
    conn.commit()
    conn.close()
    return {"ok": True, "added_checks": add}


@app.post("/api/admin/checks")
def admin_checks(data: AdminChecksIn, user=Depends(_auth)):
    _admin(user)
    conn = get_connection()
    conn.execute("UPDATE users SET checks_left = COALESCE(checks_left, 0) + ? WHERE id = ?",
                 (data.amount, data.user_id))
    conn.commit()
    conn.close()
    return {"ok": True}


@app.post("/api/admin/promo")
def admin_promo_create(data: AdminPromoIn, user=Depends(_auth)):
    _admin(user)
    conn = get_connection()
    conn.execute("CREATE TABLE IF NOT EXISTS promos (code TEXT PRIMARY KEY, percent INTEGER, "
                 "max_uses INTEGER, used_count INTEGER DEFAULT 0, disabled INTEGER DEFAULT 0, "
                 "created_at TEXT DEFAULT (datetime('now')))")
    conn.execute("INSERT OR REPLACE INTO promos (code, percent, max_uses, used_count, disabled) "
                 "VALUES (?,?,?,?,0)", (data.code.upper(), data.percent, data.max_uses))
    conn.commit()
    conn.close()
    return {"ok": True}


@app.post("/api/admin/promo_toggle")
def admin_promo_toggle(data: AdminPromoToggleIn, user=Depends(_auth)):
    _admin(user)
    conn = get_connection()
    conn.execute("UPDATE promos SET disabled = 1 - COALESCE(disabled, 0) WHERE code = ?",
                 (data.code.upper(),))
    conn.commit()
    conn.close()
    return {"ok": True}


@app.post("/api/admin/support_reply")
def admin_support_reply(data: AdminSupportReplyIn, user=Depends(_auth)):
    _admin(user)
    conn = get_connection()
    done = False
    for tbl in ("support_messages", "support"):
        try:
            cur = conn.execute(f"UPDATE {tbl} SET reply = ? WHERE id = ?",
                               (data.text, data.message_id))
            if cur.rowcount:
                done = True
                break
        except Exception:
            continue
    conn.commit()
    conn.close()
    return {"ok": done}


@app.post("/api/admin/consult_reply")
def admin_consult_reply(data: AdminConsultReplyIn, user=Depends(_auth)):
    _admin(user)
    conn = get_connection()
    done = False
    for tbl in ("consults", "consult_requests"):
        try:
            cur = conn.execute(f"UPDATE {tbl} SET reply = ? WHERE id = ?",
                               (data.text, data.id))
            if cur.rowcount:
                done = True
                break
        except Exception:
            continue
    conn.commit()
    conn.close()
    return {"ok": done}


@app.post("/api/admin/reset_password")
def admin_reset_password(data: AdminUserIdIn, user=Depends(_auth)):
    _admin(user)
    import secrets
    import string as _s
    pwd = "".join(secrets.choice(_s.ascii_letters + _s.digits) for _ in range(10))
    try:
        from utils.auth import set_password as _sp
        _sp(data.user_id, pwd)
    except Exception:
        from utils.auth import hash_password as _hp
        conn = get_connection()
        conn.execute("UPDATE users SET password_hash = ? WHERE id = ?",
                     (_hp(pwd), data.user_id))
        conn.commit()
        conn.close()
    return {"ok": True, "password": pwd}


@app.post("/api/admin/user_delete")
def admin_user_delete(data: AdminUserIdIn, user=Depends(_auth)):
    _admin(user)
    try:
        from utils.auth import delete_user
        delete_user(data.user_id)
    except Exception:
        conn = get_connection()
        conn.execute("DELETE FROM users WHERE id = ?", (data.user_id,))
        conn.commit()
        conn.close()
    return {"ok": True}